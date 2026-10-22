import io,json,tempfile,unittest
from pathlib import Path
from unittest.mock import patch
from document_review import clean_text,clean_blocks,original_context,needs_source,review_source
from lab import Lab
from recognizer import CodexRecognizer,grounded_candidates
from resume_skill import CONTRACT
from test_recognizer import field

class DocumentReviewTests(unittest.TestCase):
 def test_cleanup_preserves_facts_and_original(self):
  text='![photo](data:image/jpeg;base64...)**项目说明**\n\n- **C++**  转化率  12.5%，利润 -3 元，A/B 测试。\n\n| 奖项 | 结果 |\n| --- | --- |\n| 创新赛 | 国二 |'
  b=clean_blocks([{'text':text,'location':'Markdown'}])[0]
  self.assertEqual(b['raw_text'],text);self.assertNotIn('base64',b['text']);self.assertNotIn('**',b['text']);self.assertIn('C++ 转化率 12.5%，利润 -3 元，A/B 测试。',b['text']);self.assertIn('创新赛 国二',b['text'])
 def test_cleanup_never_guesses_date_binding(self):
  text='## 教育\n**甲校**\n**乙校**\n2021/09-2025/06\n2025/09-2027/06'
  cleaned=clean_text(text);self.assertEqual(cleaned.splitlines(),['教育','甲校','乙校','2021/09-2025/06','2025/09-2027/06'])
 def test_original_pdf_retains_position_and_renders(self):
  from reportlab.pdfgen import canvas
  stream=io.BytesIO();pdf=canvas.Canvas(stream);pdf.drawString(40,700,'Alpha University');pdf.drawString(340,700,'2021/09-2025/06');pdf.drawString(40,650,'Beta University');pdf.drawString(340,650,'2025/09-2027/06');pdf.save()
  with tempfile.TemporaryDirectory() as d:
   blocks,images,notes=original_context('fiction.pdf',stream.getvalue(),Path(d))
   alpha=next(l for l in blocks[0]['text'].splitlines() if 'Alpha' in l);self.assertIn('2021/09-2025/06',alpha);self.assertEqual(len(images),1);self.assertTrue(Path(images[0]).is_file());self.assertFalse(notes)
 def test_word_table_source_keeps_same_row(self):
  from docx import Document
  doc=Document();t=doc.add_table(rows=2,cols=2);t.cell(0,0).text='甲校';t.cell(0,1).text='2021-2025';t.cell(1,0).text='乙校';t.cell(1,1).text='2025-2027';stream=io.BytesIO();doc.save(stream)
  with tempfile.TemporaryDirectory() as d,patch('document_review.shutil.which',return_value=None):
   blocks,images,notes=original_context('fiction.docx',stream.getvalue(),Path(d));self.assertEqual(blocks[0]['text'].splitlines(),['甲校 | 2021-2025','乙校 | 2025-2027']);self.assertFalse(images);self.assertTrue(notes)
 def test_doubt_requests_original_not_generic_disclaimer(self):
  self.assertTrue(needs_source({'questions':[{'kind':'source'}]}));self.assertTrue(needs_source({'warnings':['日期与经历分离，暂不绑定。']}));self.assertFalse(needs_source({'warnings':['未提供原文件，无法核验串列。']}))
 def test_review_passes_images_and_original_line_numbers(self):
  original=[{'text':'甲校 2021/09-2025/06','location':'PDF第1页'}];result={'parser_contract':CONTRACT,'entries':[],'warnings':[],'questions':[]}
  class Engine:
   def call_with_images(self,prompt,schema,images):
    self.prompt=prompt;self.images=images;return result,{}
  engine=Engine()
  with patch('document_review.original_context',return_value=(original,['page.png'],[])):
   blocks,out,info=review_source(engine,[],[{'text':'甲校\n2021/09-2025/06','location':'Markdown'}],result,'x.pdf',b'pdf')
  self.assertEqual(len(blocks),2);self.assertTrue(info['visual']);self.assertIn('PDF第1页',engine.prompt);self.assertEqual(engine.images,['page.png'])
 def test_clean_description_grounding_still_rejects_truncation(self):
  raw='**负责** 开发。\n- 提升 20%。';entry={'category':'employment','title':'甲公司','record_type':'','source_range':{'block':0,'start_line':1,'end_line':3},'description_range':{'block':0,'start_line':2,'end_line':3},'fields':[field('company','单位名称','甲公司'),field('description','经历描述',clean_text(raw),quote=raw,kind='versioned_text')]}
  with tempfile.TemporaryDirectory() as d:
   lab=Lab(d);blocks=[{'text':'甲公司\n'+raw,'location':'原件'}];result={'parser_contract':CONTRACT,'entries':[entry]};rows,rejected=grounded_candidates(lab.state(),blocks,'x.pdf',result);self.assertFalse(rejected)
   entry['fields'][1]['value']='负责 开发。';rows,rejected=grounded_candidates(lab.state(),blocks,'x.pdf',result);self.assertTrue(rejected);lab.db.close()

 def test_cleanup_retains_urls_and_code_identifiers(self):
  result=clean_text('[项目作品](https://example.org/project)\n使用 `__init__` 与 C++，增长 +12%。')
  self.assertIn('https://example.org/project',result);self.assertIn('__init__',result);self.assertIn('C++',result);self.assertIn('+12%',result)
