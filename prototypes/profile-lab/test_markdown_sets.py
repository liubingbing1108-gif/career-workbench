import copy
import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from lab import Lab, LabError, extract_text
from recognizer import CodexRecognizer
from test_recognizer import FakeRecognizer, field


class MarkdownSetTests(unittest.TestCase):
    def test_supported_document_formats_use_real_markitdown(self):
        from docx import Document
        from pptx import Presentation
        from openpyxl import Workbook
        from reportlab.pdfgen import canvas
        documents = {'resume.txt': b'Example award', 'resume.md': b'# Example award',
                     'resume.html': b'<h1>Example award</h1>', 'resume.csv': b'award,date\nExample award,2024',
                     'resume.json': b'{"award":"Example award"}', 'resume.xml': b'<award>Example award</award>'}
        stream=io.BytesIO();doc=Document();doc.add_paragraph('Example award');doc.save(stream);documents['resume.docx']=stream.getvalue()
        stream=io.BytesIO();doc=Presentation();slide=doc.slides.add_slide(doc.slide_layouts[1]);slide.shapes.title.text='Example award';doc.save(stream);documents['resume.pptx']=stream.getvalue()
        stream=io.BytesIO();doc=Workbook();doc.active.append(['Example award']);doc.save(stream);documents['resume.xlsx']=stream.getvalue()
        stream=io.BytesIO();doc=canvas.Canvas(stream);doc.drawString(50,700,'Example award');doc.save();documents['resume.pdf']=stream.getvalue()
        for filename,raw in documents.items():
            with self.subTest(filename=filename):
                blocks=extract_text(filename,raw)
                self.assertIn('Example award',blocks[0]['text'])
                self.assertIn('MarkItDown',blocks[0]['location'])

    def test_conversion_failure_retains_original_and_never_calls_model(self):
        with tempfile.TemporaryDirectory() as directory:
            engine=FakeRecognizer([]);lab=Lab(directory,engine)
            try:
                with patch('lab.extract_text',side_effect=LabError('MarkItDown 转换失败')):
                    task=lab.upload('resume.docx',b'broken document',asynchronous=True)
                    lab.import_workers[task['id']].join(5)
                saved=lab.state()['imports'][0]
                self.assertEqual(saved['status'],'失败');self.assertFalse(saved['candidates'])
                self.assertEqual(lab.file_bytes(saved['file_id'])[1],b'broken document')
            finally:lab.db.close()

    def test_language_and_award_fields_are_atomic_independent_sets(self):
        text='English CET6 510\nFrench B2\nExample award 2024\nAnother award 2023'
        result={'entries':[
            {'category':'language','title':'English','fields':[field('language_name','语言／考试','CET6'),field('language_score','语言成绩','510')]},
            {'category':'language','title':'French','fields':[field('language_name','语言／考试','French'),field('language_score','语言成绩','B2')]},
            {'category':'honors','title':'Example award','fields':[field('award_name','荣誉名称','Example award'),field('award_date','获奖日期','2024',kind='partial_date')]},
            {'category':'honors','title':'Another award','fields':[field('award_name','荣誉名称','Another award'),field('award_date','获奖日期','2023',kind='partial_date')]}], 'warnings':[]}
        with tempfile.TemporaryDirectory() as directory:
            lab=Lab(directory,FakeRecognizer([result]))
            try:
                preview=lab.upload('resume.md',text.encode());before=copy.deepcopy(lab.state())
                self.assertEqual(preview['converter'],'MarkItDown')
                self.assertEqual(preview['document_blocks'][0]['text'],text)
                rows=preview['candidates'];self.assertEqual(len({c['group'] for c in rows}),4)
                with self.assertRaisesRegex(LabError,'完整资料组'):
                    lab.resolve_import([dict(candidate_id=rows[0]['id'],mode='new')])
                with self.assertRaisesRegex(LabError,'完整资料组'):
                    lab.resolve_import([dict(candidate_id=c['id'],mode='new' if i==0 else 'ignore') for i,c in enumerate(rows[:2])])
                self.assertEqual(lab.state(),before)
                lab.resolve_import([dict(candidate_id=c['id'],mode='new') for c in rows])
                entries=[e for e in lab.state()['entries'] if e['category'] in {'language','honors'}]
                self.assertEqual(len(entries),4)
                self.assertTrue(all(len(e['values'])==2 for e in entries))
            finally:lab.db.close()

    def test_model_receives_markdown_and_set_contract(self):
        class Capture(CodexRecognizer):
            def call(self,prompt,output_schema):
                self.prompt=prompt;return {'entries':[],'warnings':[]},{}
        engine=Capture();engine.agent_command='offline-capture'
        engine.extract([],extract_text('resume.md',b'# Example award'))
        self.assertIn('"input_format": "markdown"',engine.prompt)
        self.assertIn('同一对象',engine.prompt)
        self.assertIn('"id": "honors", "label": "荣誉与奖励", "repeatable": true',engine.prompt)


if __name__=='__main__':unittest.main()
