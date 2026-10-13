"""Generate labelled fictional samples. Requires the development python-docx/reportlab dependencies."""
from pathlib import Path
from docx import Document
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parent / "web" / "samples"
VALUES = [
    ("姓名", "Name", "Alex Example"), ("手机", "Phone", "00000000000"),
    ("邮箱", "Email", "sample@example.invalid"),
    ("学校", "School", "Example University"), ("专业", "Major", "Finance"),
    ("学历", "Degree", "Bachelor"), ("毕业日期", "Graduation date", "2027-06-30"),
    ("公司", "Company", "Example Research"), ("职位", "Role", "Finance Intern"),
    ("经历性质", "Employment type", "Internship"), ("开始日期", "Start date", "2026-06-01"),
    ("结束日期", "End date", "2026-08-31"),
    ("经历描述", "Description", "Organized source records and prepared weekly summaries."),
]

def generate():
    ROOT.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading("Fictional resume - verification only", 0)
    for label, _, value in VALUES:
        doc.add_paragraph(f"{label}：{value}")
    doc.save(ROOT / "sample.docx")
    pdf = canvas.Canvas(str(ROOT / "sample.pdf"))
    pdf.setFont("Helvetica", 11)
    pdf.drawString(50, 790, "Fictional resume - verification only")
    for i, (_, label, value) in enumerate(VALUES):
        pdf.drawString(50, 760 - i * 32, f"{label}: {value}")
    pdf.save()
    (ROOT / "sample.txt").write_text("Fictional resume - verification only\n" + "\n".join(f"{label}: {value}" for _, label, value in VALUES) + "\n")

if __name__ == "__main__":
    generate()
